# -*- coding: utf-8 -*-
"""A presença anotada na lista de turma do MGM, para o histórico do sistema.

    set -a && . ../.secrets/4yu.env && set +a
    python3 planilhas/importa_presencas.py <planilha.xlsx>... [--financeiro CONTROLE.xlsx] [--pula AAAA-MM-DD] [--confirmo]

Pedido do MGM (07/out/2026): o histórico desde janeiro de 2025, que só existia
nas planilhas mensais. Legenda decidida com o Gabriel nesta data:

- `P...` (P, P REP, P ANT, P EXP...): presente. Com "REP", origem reposição, e
  "REP dd/mm" liga a reposição à falta daquele dia.
- `F`, `F EXP`, `F REP`, `F ANT`: falta. `FAR`: falta justificada. `LIC`: licença.
- `FER`, `X`, `XX`, horário solto, `FISIO`, `FASCIA` e o resto: ignorados.

Quem é quem:

- Pela matrícula da linha. Matrícula que o sistema não conhece é ex-aluno:
  entra como pessoa **inativa**, com nome e número, para o histórico não sumir.
  O nome vem do controle de pagamentos (`--financeiro`), que tem o nome
  completo de todo aluno que já passou pelo estúdio, ativo ou inativo; e quem
  só está lá também concorre no casamento por nome.
- Sem matrícula, pelo nome, só quando aponta **um** aluno. Nome de uma palavra
  ("MARIA") desempata pelo horário: fica quem frequenta aquele dia e hora (pela
  matrícula em outro mês do lote ou pela vaga atual).

O que não muda: registro feito no sistema (recepção, bot, profissional) nunca
é sobrescrito; data futura e `--pula` ficam de fora; linha de personal,
domicílio e fisioterapia também.

Aula que não existe é criada: da série, quando o horário está na grade de
hoje; senão como aula de pilates sem série, porque a grade de 2025 não é a
de agora. Pode rodar de novo: o que já está gravado não duplica.
"""
import json
import os
import re
import sys
import unicodedata
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402
from importa_grade import DIA_NUM, limpo, norm_hora, sem_desenho, sql  # noqa: E402
from importa_vagas import FORA, MES, SERVICO, SLUG  # noqa: E402


def chave(nome):
    """Nome sem acento, caixa e espaço sobrando: "Lucília  Cini" = "LUCILIA CINI"."""
    s = unicodedata.normalize('NFKD', nome).encode('ascii', 'ignore').decode().upper()
    return re.sub(r'\s+', ' ', re.sub(r'\(.*?\)', '', s)).strip()


def status_de(celula):
    """(status, é_reposição) ou None quando a célula não vira registro."""
    v = re.sub(r'\s+', ' ', celula.strip().upper())
    if not v:
        return None
    rep = 'REP' in v
    if v.startswith('P') and not v.startswith('PERS'):
        return ('presente', rep)
    if v == 'FAR':
        return ('falta_avisada', False)
    if v == 'F' or re.match(r'^F (EXP|REP|ANT)', v):
        return ('falta', rep)
    if v == 'LIC':
        return ('licenca', False)
    return None


# A resposta do cliente para os nomes soltos da planilha (só o primeiro nome):
# quem ainda é aluno, pela matrícula, e onde ele disse o horário. Mora fora do
# git, em `.rascunho/respostas-presenca.json`, porque tem nome de aluno e este
# repositório é público. Formato:
#   {"resposta": {"MARIA": ["001", "002"]}, "no_horario": [["MARIA", "SEXTA", "17:00", "002"]]}
# Nome com mais de um candidato vai para quem aparece naquele dia e hora; nome
# que não está lá segue as regras de sempre.
def _respostas():
    caminho = os.path.join(os.path.dirname(__file__), '..', '.rascunho', 'respostas-presenca.json')
    if not os.path.exists(caminho):
        return {}, {}
    with open(caminho, encoding='utf-8') as f:
        r = json.load(f)
    return r.get('resposta', {}), {(n, d, h): m for n, d, h, m in r.get('no_horario', [])}


RESPOSTA, RESPOSTA_NO_HORARIO = _respostas()

# apelido -> nome do cadastro: "ZÉ MARIA" é o José Maria
APELIDO = {'ZE': 'JOSE', 'OSWALDINHO': 'OSWALDO', 'CRIS': 'CRISTINA', 'CAROL': 'CAROLINA', 'GABI': 'GABRIELA',
           'DANI': 'DANIELA', 'FER': 'FERNANDA', 'BIA': 'BEATRIZ', 'MALU': 'MARIA LUIZA'}


def parece(parte, palavra):
    """Começo da palavra, ou a mesma com um erro de digitação (MASSERAM/MASSERAN)."""
    if palavra.startswith(parte):
        return True
    if len(parte) < 5:
        return False
    from difflib import SequenceMatcher
    return SequenceMatcher(None, parte, palavra[:len(parte) + 1]).ratio() >= 0.8


def abrevia(partes, nome):
    """As palavras da planilha começam palavras do nome, na ordem, a 1ª na 1ª."""
    partes = ' '.join(APELIDO.get(x, x) for x in partes).split(' ')
    if not partes or not nome:
        return False
    # nome solto ("PAULA") só pelo começo exato: com folga, PAULA vira PAULO.
    # Com sobrenome junto, o primeiro nome pode ter erro (THAYS YANO)
    if not (nome[0].startswith(partes[0]) or (len(partes) > 1 and parece(partes[0], nome[0]))):
        return False
    j = 1
    for parte in partes[1:]:
        while j < len(nome) and not parece(parte, nome[j]):
            j += 1
        if j == len(nome):
            return False
        j += 1
    return True


def falta_de(celula, ano, mes):
    """ "P REP 03/09" -> a data da falta que esta reposição paga, ou None."""
    m = re.search(r'REP\.?\s*(\d{1,2})/(\d{1,2})', celula.upper())
    if not m:
        return None
    d, mm = int(m.group(1)), int(m.group(2))
    try:
        return date(ano if mm <= mes else ano - 1, mm, d).isoformat()
    except ValueError:
        return None


def mes_do_arquivo(caminho):
    m = re.search(r'-\s*(\w+)\s+(\d{2,4})', os.path.basename(caminho).upper())
    if not m or m.group(1) not in MES:
        sys.exit(f'o nome do arquivo não diz o mês: {caminho}')
    a = int(m.group(2))
    return (a if a > 100 else 2000 + a), MES[m.group(1)]


def le(caminho, ano, mes):
    """Uma linha por (data, hora, aluno) com o que a célula diz."""
    wb = openpyxl.load_workbook(sem_desenho(caminho), data_only=True)
    registros, ignorados = [], {}
    for ws in wb.worksheets:
        dia = ws.title.strip().upper()
        if dia not in DIA_NUM:
            continue
        cabecas = [r for r in range(1, ws.max_row + 1) if limpo(ws.cell(r, 1).value) == 'Vagas']
        for i, cabeca in enumerate(cabecas):
            fim = cabecas[i + 1] - 1 if i + 1 < len(cabecas) else ws.max_row
            hora = norm_hora(ws.cell(cabeca + 1, 2).value)
            if not hora:
                continue
            colunas = {}
            for c in range(9, ws.max_column + 1):
                m = re.match(r'\s*(\d{1,2})', limpo(ws.cell(cabeca, c).value))
                if m:
                    try:
                        colunas[c] = date(ano, mes, int(m.group(1))).isoformat()
                    except ValueError:
                        pass
            for r in range(cabeca + 1, fim + 1):
                mat = limpo(ws.cell(r, 4).value)
                nome = re.sub(r'\s+', ' ', limpo(ws.cell(r, 5).value))
                if not nome or FORA.search(nome) or 'N O M E' in nome:
                    continue
                na_vaga = isinstance(ws.cell(r, 1).value, int)
                for c, dia_iso in colunas.items():
                    celula = limpo(ws.cell(r, c).value)
                    st = status_de(celula)
                    if not st:
                        if celula:
                            ignorados[celula.upper()] = ignorados.get(celula.upper(), 0) + 1
                        continue
                    registros.append({'data': dia_iso, 'dia': dia, 'hora': hora,
                                      'mat': str(int(mat)) if mat.isdigit() else '',
                                      'nome': nome, 'status': st[0], 'rep': st[1],
                                      'falta': falta_de(celula, ano, mes) if st[1] else None,
                                      'na_vaga': na_vaga})
    return registros, ignorados


def le_financeiro(caminho):
    """matrícula -> nome completo, das abas ATIVOS e INATIVOS do controle."""
    from importa_contratos import sem_desenho as sem_desenho_pag, so_digito
    wb = openpyxl.load_workbook(sem_desenho_pag(caminho), data_only=True, read_only=True)
    alunos = {}
    for ws in wb.worksheets:
        for linha in ws.iter_rows(values_only=True):
            nome = re.sub(r'\s+', ' ', str(linha[0]).strip()) if linha and linha[0] else ''
            mat = so_digito(linha[1]) if len(linha) > 1 else ''
            if nome and mat:
                alunos[str(int(mat))] = nome
    return alunos


def main():
    if not os.environ.get('SUPABASE_ACCESS_TOKEN'):
        sys.exit('falta SUPABASE_ACCESS_TOKEN: carregue o .secrets/4yu.env')
    seco = '--confirmo' not in sys.argv
    pula = {sys.argv[i + 1] for i, a in enumerate(sys.argv) if a == '--pula'}
    fin = [sys.argv[i + 1] for i, a in enumerate(sys.argv) if a == '--financeiro']
    arquivos = [a for a in sys.argv[1:] if not a.startswith('--') and a not in pula and a not in fin]
    roster = le_financeiro(fin[0]) if fin else {}
    if not arquivos:
        sys.exit('uso: importa_presencas.py <planilha.xlsx>... [--pula AAAA-MM-DD] [--confirmo]')

    conta = sql(f"select id, fuso from app_verandi.conta where slug = '{SLUG}'")[0]
    conta_id, fuso = conta['id'], conta['fuso']
    hoje = sql(f"select (now() at time zone '{fuso}')::date::text as d")[0]['d']
    servico_id = sql(f"select id from app_verandi.servico where conta_id = '{conta_id}' "
                     f"and nome = '{SERVICO}'")[0]['id']

    lidos, ignorados = [], {}
    for arq in arquivos:
        ano, mes = mes_do_arquivo(arq)
        regs, ign = le(arq, ano, mes)
        lidos += regs
        for k, n in ign.items():
            ignorados[k] = ignorados.get(k, 0) + n

    def pessoas():
        t = sql("select id, identificador_externo as m, nome from app_verandi.pessoa "
                f"where conta_id = '{conta_id}'")
        return t, {str(int(p['m'])): p for p in t if (p['m'] or '').isdigit()}
    todas, por_mat = pessoas()

    # quem o sistema não conhece: do financeiro (nome completo) e as matrículas
    # que só as planilhas trazem. Viram candidatos `novo:<matrícula>`; só os
    # que casarem com alguma presença são cadastrados, inativos
    # quem está no financeiro com outra grafia do mesmo nome ("GUSTAVO SOUZA
    # SILVA" e "GUSTAVO SOUZA DA SILVA") já é a pessoa do sistema: não concorre
    from difflib import SequenceMatcher
    nomes_sistema = [chave(p['nome']) for p in todas]
    ex = {m: n for m, n in roster.items() if m not in por_mat
          and not any(SequenceMatcher(None, chave(n), x).ratio() >= 0.88 for x in nomes_sistema)}
    for r in lidos:
        if r['mat'] and r['mat'] not in por_mat and r['mat'] not in roster \
                and len(r['nome']) > len(ex.get(r['mat'], '')):
            ex[r['mat']] = r['nome']
    for m, n in ex.items():
        por_mat[m] = {'id': f'novo:{m}', 'm': m, 'nome': n}

    series = {(s['dia'], s['hora']): s for s in sql(
        "select id, dia_semana as dia, to_char(hora_inicio, 'HH24:MI') as hora "
        f"from app_verandi.serie where conta_id = '{conta_id}' and servico_id = '{servico_id}'")}
    do_horario = {}
    for v in sql("select v.pessoa_id, s.dia_semana as dia, to_char(s.hora_inicio, 'HH24:MI') as hora "
                 "from app_verandi.vaga v join app_verandi.serie s on s.id = v.serie_id "
                 f"where v.conta_id = '{conta_id}'"):
        do_horario.setdefault((v['dia'], v['hora']), set()).add(v['pessoa_id'])
    ativos_no_mes = {}
    for r in lidos:
        if r['mat'] in por_mat:
            do_horario.setdefault((DIA_NUM[r['dia']], r['hora']), set()).add(por_mat[r['mat']]['id'])
            ativos_no_mes.setdefault(r['data'][:7], set()).add(por_mat[r['mat']]['id'])

    candidatos = todas + [p for p in por_mat.values() if str(p['id']).startswith('novo:')]

    no_mes = {}
    for r in lidos:
        if r['mat'] in por_mat:
            no_mes.setdefault((DIA_NUM[r['dia']], r['hora'], r['data'][:7]), set()).add(por_mat[r['mat']]['id'])

    # quantas vezes cada um aparece com nome completo em cada dia e hora
    vezes_no_horario = {}
    for r in lidos:
        if r['mat'] in por_mat:
            k = (DIA_NUM[r['dia']], r['hora'], por_mat[r['mat']]['id'])
            vezes_no_horario[k] = vezes_no_horario.get(k, 0) + 1

    def respondido(nome, dia, hora, mes):
        """Um candidato só: é ele, foi o cliente quem disse. Mais de um: quem
        aparece naquele dia e hora no mesmo mês; senão, quem tem a vaga fixa;
        senão, quem mais frequenta aquele horário. Empate fica de fora."""
        fixo = RESPOSTA_NO_HORARIO.get((chave(nome), dia, hora))
        if fixo in por_mat:
            return por_mat[fixo]
        pool = [por_mat[m] for m in RESPOSTA.get(chave(nome), []) if m in por_mat]
        if len(pool) == 1:
            return pool[0]
        for ids in (no_mes.get((DIA_NUM[dia], hora, mes), set()),
                    do_horario.get((DIA_NUM[dia], hora), set())):
            daqui = [p for p in pool if p['id'] in ids]
            if len(daqui) == 1:
                return daqui[0]
        contagem = sorted(((vezes_no_horario.get((DIA_NUM[dia], hora, p['id']), 0), p) for p in pool),
                          key=lambda x: -x[0])
        if contagem and contagem[0][0] > 0 and (len(contagem) == 1 or contagem[0][0] > contagem[1][0]):
            return contagem[0][1]
        return None

    def por_nome(nome, dia, hora, mes):
        """Primeiro entre quem já está no sistema; só sem resposta, o financeiro."""
        if chave(nome) in RESPOSTA:
            return respondido(nome, dia, hora, mes)
        return (casa(nome, dia, hora, todas) or casa(nome, dia, hora, candidatos)
                or desempata(nome, dia, hora, mes, todas) or desempata(nome, dia, hora, mes, candidatos))

    def desempata(nome, dia, hora, mes, candidatos):
        """Nome solto ou com erro: candidatos pelo começo do nome, apelido ou
        grafia parecida; fica quem frequenta o horário e, empatando, quem estava
        ativo no mês."""
        partes = ' '.join(APELIDO.get(x, x) for x in chave(nome).split(' ')).split(' ')
        if not partes or len(partes[0]) < 3:
            return None
        pool = [p for p in candidatos
                if (lambda n: parece(partes[0], n[0]) and all(
                    any(parece(x, w) for w in n[1:]) for x in partes[1:]))(chave(p['nome']).split(' '))]
        if not pool:
            return None
        daqui = [p for p in pool if p['id'] in do_horario.get((DIA_NUM[dia], hora), set())]
        if len(daqui) == 1:
            return daqui[0]
        ativos = [p for p in (daqui or pool) if p['id'] in ativos_no_mes.get(mes, set())]
        if daqui and len(ativos) == 1:
            return ativos[0]
        return None

    def casa(nome, dia, hora, candidatos):
        k = chave(nome)
        if len(k) < 3:
            return None
        iguais = [p for p in candidatos if chave(p['nome']) == k]
        if len(iguais) == 1:
            return iguais[0]
        comeca = [p for p in candidatos if chave(p['nome']).startswith(k + ' ')] or iguais
        if len(comeca) == 1:
            return comeca[0]
        if not comeca and len(k.split(' ')) == 1:
            # nome solto com erro de digitação, quando só um nome do cadastro
            # chega perto (CLAUDINIEI -> Claudinei)
            from difflib import SequenceMatcher
            perto = [p for p in candidatos
                     if SequenceMatcher(None, k, chave(p['nome']).split(' ')[0]).ratio() >= 0.88]
            if len(perto) == 1:
                return perto[0]
        if not comeca and len(k.split(' ')) > 1:
            # o nome da planilha é o do meio: "MARCELO SOLER" é Marcos Marcelo Soler
            meio = [p for p in candidatos
                    if all(any(w.startswith(x) for w in chave(p['nome']).split(' ')) for x in k.split(' '))]
            if len(meio) == 1:
                return meio[0]
        if not comeca:
            # abreviado: "THAIS YANO", "CRIS LUCENA", "THAIS Y". Cada palavra da
            # planilha começa uma palavra do nome, na ordem, e a primeira é o
            # começo do primeiro nome
            partes = k.split(' ')
            comeca = [p for p in candidatos if abrevia(partes, chave(p['nome']).split(' '))]
            if len(comeca) == 1:
                return comeca[0]
        daqui = [p for p in comeca if p['id'] in do_horario.get((DIA_NUM[dia], hora), set())]
        return daqui[0] if len(daqui) == 1 else None

    unicos, futuros, pulados, sem_pessoa, n_nome = {}, 0, 0, {}, 0
    for r in lidos:
        if r['data'] in pula:
            pulados += 1
            continue
        if r['data'] > hoje:
            futuros += 1
            continue
        pessoa = por_mat.get(r['mat']) if r['mat'] else None
        if not pessoa:
            pessoa = por_nome(r['nome'], r['dia'], r['hora'], r['data'][:7])
            n_nome += bool(pessoa)
        if not pessoa:
            sem_pessoa[r['nome']] = sem_pessoa.get(r['nome'], 0) + 1
            continue
        serie = series.get((DIA_NUM[r['dia']], r['hora']))
        origem = 'reposicao' if r['rep'] else ('recorrente' if r['na_vaga'] and serie else 'avulso')
        unicos[(r['data'], r['hora'], pessoa['id'])] = (
            serie['id'] if serie else None, r['status'], origem, r['falta'])

    # horário de costume: onde a pessoa aparece 2+ vezes no mês. Aparecer fora
    # dele, sem vaga ali, é reposição (ou adiantamento), não aula avulsa
    costume = {}
    for (d, h, p) in unicos:
        dow = date.fromisoformat(d).isoweekday() % 7
        chave_c = (p, d[:7], dow, h)
        costume[chave_c] = costume.get(chave_c, 0) + 1
    n_rep = 0
    for (d, h, p), (s_id, st, o, f) in list(unicos.items()):
        dow = date.fromisoformat(d).isoweekday() % 7
        if o == 'avulso' and costume.get((p, d[:7], dow, h), 0) < 2 \
                and p not in do_horario.get((dow, h), set()):
            unicos[(d, h, p)] = (s_id, st, 'reposicao', f)
            n_rep += 1
    print(f'  fora do horário de costume, como reposição: {n_rep}')

    por_status, por_mes = {}, {}
    for (d, _, _), (_, st, _, _) in unicos.items():
        por_status[st] = por_status.get(st, 0) + 1
        por_mes[d[:7]] = por_mes.get(d[:7], 0) + 1
    usados = sorted({p[5:] for (_, _, p) in unicos if str(p).startswith('novo:')})
    print(f'{len(arquivos)} planilhas: {len(unicos)} registros até {hoje}; '
          f'{n_nome} casados pelo nome; {len(usados)} ex-alunos a cadastrar inativos '
          f'({sum(1 for m in usados if m in roster)} com nome do financeiro)')
    print('  ' + ', '.join(f'{k}: {v}' for k, v in sorted(por_status.items())))
    print('  por mês: ' + ', '.join(f'{k} {v}' for k, v in sorted(por_mes.items())))
    print(f'  fora: {futuros} futuros, {pulados} do dia pulado; '
          f"{sum(1 for v in unicos.values() if v[0] is None)} em horário fora da grade atual")
    print(f'  reposições com data da falta: {sum(1 for v in unicos.values() if v[3])}')
    if sem_pessoa:
        print(f'  sem cadastro ou ambíguo ({len(sem_pessoa)} nomes, {sum(sem_pessoa.values())} registros): '
              + '; '.join(f'{k} ({v})' for k, v in sorted(sem_pessoa.items(), key=lambda x: -x[1])[:50]))

    if seco:
        print('  ensaio: nada foi gravado. Para valer: --confirmo')
        return

    def cita(v):
        return 'null' if v is None else f"'{v}'"

    if usados:
        valores = ',\n'.join(
            f"('{conta_id}', '{ex[m].replace(chr(39), chr(39) * 2)}', '{m.zfill(3)}', false)" for m in usados)
        sql('insert into app_verandi.pessoa (conta_id, nome, identificador_externo, ativo) '
            f'values {valores};')
        _, reais = pessoas()
        unicos = {(d, h, reais[p[5:]]['id'] if str(p).startswith('novo:') else p): v
                  for (d, h, p), v in unicos.items()}

    itens = list(unicos.items())
    antes = sql(f"select count(*)::int as n from app_verandi.participacao where conta_id = '{conta_id}'")[0]['n']
    for i in range(0, len(itens), 1500):
        lote = itens[i:i + 1500]
        linhas = ',\n'.join(
            f"({cita(s)}::uuid, '{d}'::date, '{h}'::time, '{p}'::uuid, '{st}', '{o}', {cita(f)}::date)"
            for (d, h, p), (s, st, o, f) in lote)
        sql(f"""
begin;
create temp table imp (serie_id uuid, dia date, hora time, pessoa_id uuid, status text,
                       origem text, falta date) on commit drop;
insert into imp values
{linhas};

insert into app_verandi.sessao (conta_id, serie_id, servico_id, profissional_id, local_id,
                                inicio, duracao_min, capacidade, status)
select distinct se.conta_id, se.id, se.servico_id, se.profissional_id, se.local_id,
       (i.dia + i.hora) at time zone '{fuso}', se.duracao_min, se.capacidade,
       'prevista'::app_verandi.status_sessao
from imp i join app_verandi.serie se on se.id = i.serie_id
on conflict (serie_id, inicio) do nothing;

insert into app_verandi.sessao (conta_id, serie_id, servico_id, inicio, duracao_min,
                                capacidade, status)
select '{conta_id}', null, '{servico_id}', x.inicio, 60, 4, 'prevista'::app_verandi.status_sessao
from (select distinct (dia + hora) at time zone '{fuso}' as inicio from imp where serie_id is null) x
where not exists (select 1 from app_verandi.sessao s where s.conta_id = '{conta_id}'
                  and s.serie_id is null and s.servico_id = '{servico_id}' and s.inicio = x.inicio);

insert into app_verandi.participacao (conta_id, sessao_id, pessoa_id, origem, status,
                                      registrado_por_origem)
select '{conta_id}', s.id, i.pessoa_id, i.origem::app_verandi.origem_participacao,
       i.status::app_verandi.status_participacao, 'importacao'
from imp i
join app_verandi.sessao s on s.conta_id = '{conta_id}'
 and s.inicio = (i.dia + i.hora) at time zone '{fuso}'
 and (s.serie_id = i.serie_id or (i.serie_id is null and s.serie_id is null
      and s.servico_id = '{servico_id}'))
on conflict (sessao_id, pessoa_id) do update set status = excluded.status
where app_verandi.participacao.status in ('esperada', 'confirmada')
  and app_verandi.participacao.registrado_por_origem in ('importacao', 'sistema');

update app_verandi.participacao r set reposicao_de_id = (
  select f.id from app_verandi.participacao f join app_verandi.sessao fs on fs.id = f.sessao_id
  where f.pessoa_id = r.pessoa_id and f.status in ('falta', 'falta_avisada')
    and (fs.inicio at time zone '{fuso}')::date = i.falta
    and not exists (select 1 from app_verandi.participacao o where o.reposicao_de_id = f.id)
  order by fs.inicio limit 1)
from imp i, app_verandi.sessao s
where i.falta is not null and r.sessao_id = s.id and r.pessoa_id = i.pessoa_id
  and s.inicio = (i.dia + i.hora) at time zone '{fuso}'
  and r.reposicao_de_id is null and r.registrado_por_origem = 'importacao';
commit;
""")
        print(f'  lote {i // 1500 + 1}: {len(lote)} gravados')
    depois = sql(f"select count(*)::int as n from app_verandi.participacao where conta_id = '{conta_id}'")[0]['n']
    print(f'  participações novas: {depois - antes}')


if __name__ == '__main__':
    main()
