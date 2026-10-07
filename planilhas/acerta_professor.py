# -*- coding: utf-8 -*-
"""Quem deu cada aula do histórico do MGM, conforme a planilha do mês.

    set -a && . ../.secrets/4yu.env && set +a
    python3 planilhas/acerta_professor.py <planilha.xlsx>... [--pula AAAA-MM-DD] [--confirmo]

`importa_presencas.py` criou as aulas antigas com o professor de **hoje** da
turma, e a grade mudou de mãos ao longo de 2025: 1.025 aulas ficaram com o
professor errado, e o relatório de aulas por professor (base do pagamento
da equipe) contaria errado. A planilha escreve "Prof. MÁRCIA" acima de cada
bloco; aqui o primeiro nome casa com o cadastro da equipe.

Professor que já saiu (a Bruna, em 2025) entra como profissional inativo. Só
aula passada que veio da importação muda; aula marcada no sistema e os dias
de `--pula` ficam como estão.
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402
from importa_grade import DIA_NUM, limpo, norm_hora, sem_desenho, sql  # noqa: E402
from importa_presencas import chave, mes_do_arquivo  # noqa: E402
from importa_vagas import SLUG  # noqa: E402


def blocos(caminho):
    """(dia da semana, hora) -> primeiro nome do professor, de um mês."""
    wb = openpyxl.load_workbook(sem_desenho(caminho), data_only=True)
    saida = {}
    for ws in wb.worksheets:
        dia = ws.title.strip().upper()
        if dia not in DIA_NUM:
            continue
        for h in [r for r in range(1, ws.max_row + 1) if limpo(ws.cell(r, 1).value) == 'Vagas']:
            prof = ''
            for rr in range(max(1, h - 3), h):
                v = limpo(ws.cell(rr, 1).value)
                if v.lower().startswith('prof'):
                    prof = re.sub(r'^prof\.?\s*', '', v, flags=re.I).strip()
            hora = norm_hora(ws.cell(h + 1, 2).value)
            if prof and hora:
                saida[(DIA_NUM[dia], hora)] = chave(prof).split(' ')[0]
    return saida


def main():
    seco = '--confirmo' not in sys.argv
    pula = {sys.argv[i + 1] for i, a in enumerate(sys.argv) if a == '--pula'}
    arquivos = [a for a in sys.argv[1:] if not a.startswith('--') and a not in pula]
    conta = sql(f"select id, fuso from app_verandi.conta where slug = '{SLUG}'")[0]
    conta_id, fuso = conta['id'], conta['fuso']

    por_mes = {mes_do_arquivo(a): blocos(a) for a in arquivos}
    primeiros = {p for m in por_mes.values() for p in m.values()}

    def equipe():
        return sql(f"select id, nome from app_verandi.profissional where conta_id = '{conta_id}'")
    eq = equipe()
    faltam = sorted(n for n in primeiros
                    if not [p for p in eq if chave(p['nome']).split(' ')[0].startswith(n)])
    print(f'professores nas planilhas: {sorted(primeiros)}; sem cadastro: {faltam}')
    if faltam and not seco:
        sql('insert into app_verandi.profissional (conta_id, nome, ativo) values '
            + ', '.join(f"('{conta_id}', '{n.title()}', false)" for n in faltam) + ';')
        eq = equipe()
    id_de = {}
    for n in primeiros:
        c = [p for p in eq if chave(p['nome']).split(' ')[0].startswith(n)]
        if len(c) == 1:
            id_de[n] = c[0]['id']
        elif not seco or n not in faltam:
            print(f'  professor ambíguo ou ausente: {n}')

    linhas = []
    for (ano, mes), m in por_mes.items():
        for (dow, hora), n in m.items():
            if n in id_de:
                linhas.append(f"({ano}, {mes}, {dow}, '{hora}', '{id_de[n]}'::uuid)")
            elif seco and n in faltam:
                linhas.append(f"({ano}, {mes}, {dow}, '{hora}', null::uuid)")
    filtro = f"""
from app_verandi.sessao s, (values {', '.join(linhas)}) as v(ano, mes, dow, hora, prof)
where s.conta_id = '{conta_id}' and s.inicio < now()
  and extract(year from s.inicio at time zone '{fuso}') = v.ano
  and extract(month from s.inicio at time zone '{fuso}') = v.mes
  and extract(dow from s.inicio at time zone '{fuso}') = v.dow
  and to_char(s.inicio at time zone '{fuso}', 'HH24:MI') = v.hora
  and (s.inicio at time zone '{fuso}')::date not in ({', '.join(f"'{d}'" for d in pula) or "'1900-01-01'"})
  and s.profissional_id is distinct from v.prof
  and exists (select 1 from app_verandi.participacao x where x.sessao_id = s.id
              and x.registrado_por_origem = 'importacao')"""
    n = sql(f'select count(*)::int as n {filtro}')[0]['n']
    print(f'aulas a acertar: {n}')
    if seco:
        print('ensaio: nada foi gravado. Para valer: --confirmo')
        return
    sql(f'update app_verandi.sessao a set profissional_id = t.prof '
        f'from (select s.id, v.prof {filtro}) t where a.id = t.id;')
    print(f'acertadas: {n}')


if __name__ == '__main__':
    main()
