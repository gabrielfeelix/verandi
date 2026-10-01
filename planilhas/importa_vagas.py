# -*- coding: utf-8 -*-
"""Põe cada aluno do MGM na turma dele, lido da lista de turma do mês.

    set -a && . ../.secrets/4yu.env && set +a
    python3 planilhas/importa_vagas.py "LISTA DE TURMA E PRESENÇA - OUTUBRO 26.xlsx"
    python3 planilhas/importa_vagas.py "...xlsx" --confirmo

A grade (`importa_grade.py`) criou as séries e os alunos (`importa_alunos.py`)
criaram as pessoas, mas ninguém tinha sido posto em turma nenhuma: a chamada
saía vazia. Este script cria a **vaga**, a matrícula de uma pessoa numa série,
para cada linha da planilha que tem número de matrícula.

O que fica de fora, de propósito:

- **Personal, domicílio e fisioterapia.** A linha está escrita no bloco do
  horário porque o horário é o mesmo, mas é outro serviço e outra agenda.
- **Reposição e reserva.** Quem está ali por um mês não é dono da vaga.
- **Turma fechada.** A série inativa não gera sessão, e a vaga não teria onde
  aparecer.

A vaga começa no primeiro dia do mês da planilha e liga no contrato de pilates
vigente da pessoa, quando ele existe, para que encerrar o contrato feche a vaga.
As sessões já geradas daquele dia em diante ganham a participação agora; as
que ainda não existem nascem com ela quando alguém abrir a agenda.

A presença anotada na folha não entra: as letras (P, F, FAR, X, LIC) não têm
legenda, e chamada errada vira reposição errada.

Roda uma vez: se a conta já tiver vaga, para e diz.
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402
from importa_grade import DIA_NUM, limpo, norm_hora, sem_desenho, sql  # noqa: E402

SLUG = 'mgm-pilates'
SERVICO = 'Pilates aparelho'
FORA = re.compile(r'PERS|PERSONAL|DOMIC|FISIO|REPOSI|RESERVA', re.I)
MES = {'JANEIRO': 1, 'FEVEREIRO': 2, 'MARÇO': 3, 'ABRIL': 4, 'MAIO': 5, 'JUNHO': 6, 'JULHO': 7,
       'AGOSTO': 8, 'SETEMBRO': 9, 'OUTUBRO': 10, 'NOVEMBRO': 11, 'DEZEMBRO': 12}


def le_alunos(caminho):
    """Uma linha por aluno em cada bloco, com o horário do bloco."""
    wb = openpyxl.load_workbook(sem_desenho(caminho), data_only=True)
    linhas = []
    for ws in wb.worksheets:
        dia = ws.title.strip().upper()
        if dia not in DIA_NUM:
            continue
        cabecas = [r for r in range(1, ws.max_row + 1) if limpo(ws.cell(r, 1).value) == 'Vagas']
        for i, cabeca in enumerate(cabecas):
            fim = cabecas[i + 1] - 1 if i + 1 < len(cabecas) else ws.max_row
            hora = norm_hora(ws.cell(cabeca + 1, 2).value)
            for r in range(cabeca + 1, fim + 1):
                mat = limpo(ws.cell(r, 4).value)
                nome = re.sub(r'\s+', ' ', limpo(ws.cell(r, 5).value))
                linha = {'dia': dia, 'hora': hora, 'mat': mat, 'nome': nome}
                if mat.isdigit():
                    linhas.append(linha)
    return linhas


def main():
    if not os.environ.get('SUPABASE_ACCESS_TOKEN'):
        sys.exit('falta SUPABASE_ACCESS_TOKEN — carregue o .secrets/4yu.env')
    if len(sys.argv) < 2:
        sys.exit('uso: python3 planilhas/importa_vagas.py <planilha.xlsx> [--confirmo]')
    seco = '--confirmo' not in sys.argv
    caminho = sys.argv[1]

    m = re.search(r'-\s*(\w+)\s+(\d{2})', os.path.basename(caminho).upper())
    if not m or m.group(1) not in MES:
        sys.exit('o nome do arquivo não diz o mês ("... - OUTUBRO 26.xlsx")')
    inicio = f'20{m.group(2)}-{MES[m.group(1)]:02d}-01'

    conta_id = sql(f"select id from app_verandi.conta where slug = '{SLUG}'")[0]['id']
    servico_id = sql(f"select id from app_verandi.servico where conta_id = '{conta_id}' "
                     f"and nome = '{SERVICO}'")[0]['id']

    jaTem = sql(f"select count(*)::int as n from app_verandi.vaga where conta_id = '{conta_id}'")[0]['n']
    if jaTem:
        sys.exit(f'a conta já tem {jaTem} vagas. Este script é para a primeira carga.')

    # matrícula sem zero à esquerda: a planilha escreve 76, o cadastro 076
    pessoas = {str(int(p['m'])): p for p in sql(
        "select id, identificador_externo as m, nome from app_verandi.pessoa "
        f"where conta_id = '{conta_id}' and ativo and identificador_externo ~ '^[0-9]+$'")}
    series = {(s['dia'], s['hora']): s for s in sql(
        "select id, dia_semana as dia, to_char(hora_inicio, 'HH24:MI') as hora, capacidade, ativo "
        f"from app_verandi.serie where conta_id = '{conta_id}' and servico_id = '{servico_id}'")}
    contratos = {c['pessoa_id']: c['id'] for c in sql(
        "select distinct on (c.pessoa_id) c.pessoa_id, c.id from app_verandi.contrato c "
        "join app_verandi.plano p on p.id = c.plano_id "
        f"where c.conta_id = '{conta_id}' and c.status <> 'encerrado' and p.servico_id = '{servico_id}' "
        "order by c.pessoa_id, c.inicio desc")}

    vagas, fora, sem_pessoa, sem_turma = {}, [], [], []
    for a in le_alunos(caminho):
        rotulo = f"{a['dia'][:3].title()} {a['hora']}  {a['mat']:>4} {a['nome']}"
        if FORA.search(a['nome']):
            fora.append(rotulo)
            continue
        pessoa = pessoas.get(str(int(a['mat'])))
        if not pessoa:
            sem_pessoa.append(rotulo)
            continue
        serie = series.get((DIA_NUM[a['dia']], a['hora']))
        if not serie or not serie['ativo']:
            sem_turma.append(rotulo)
            continue
        vagas[(serie['id'], pessoa['id'])] = (serie, pessoa, a)

    ocupacao = {}
    for serie, _, _ in vagas.values():
        ocupacao[serie['id']] = ocupacao.get(serie['id'], 0) + 1
    lotadas = [s for s in series.values() if ocupacao.get(s['id'], 0) > s['capacidade']]

    sem_contrato = sorted({p['nome'] for _, p, _ in vagas.values() if p['id'] not in contratos})
    print(f'{len(vagas)} vagas em {len(ocupacao)} turmas, de {inicio} em diante')
    print(f'{len({p["id"] for _, p, _ in vagas.values()})} alunos; '
          f'{len(sem_contrato)} sem contrato de pilates no sistema')
    for titulo, lista in (('fora (personal, reposição, reserva)', fora),
                          ('matrícula sem cadastro', sem_pessoa),
                          ('turma fechada ou inexistente', sem_turma),
                          ('sem contrato de pilates', sem_contrato)):
        if lista:
            print(f'\n  {titulo}:')
            for linha in lista:
                print(f'    {linha}')
    if lotadas:
        print('\n  turma com mais aluno fixo que lugar (a capacidade sobe):')
        for s in lotadas:
            print(f"    dia {s['dia']} {s['hora']}: {ocupacao[s['id']]}/{s['capacidade']}")

    if seco:
        print('\nensaio: nada foi gravado. Para valer: --confirmo')
        return

    def cita(v):
        return 'null' if v is None else f"'{v}'"

    comandos = ['begin;']
    for s in lotadas:
        n = ocupacao[s['id']]
        comandos.append(f"update app_verandi.serie set capacidade = {n} where id = '{s['id']}';")
        comandos.append(f"update app_verandi.sessao set capacidade = {n} where serie_id = '{s['id']}' "
                        f"and capacidade < {n} and inicio >= '{inicio}'::date;")
    for serie, pessoa, _ in vagas.values():
        comandos.append(
            'insert into app_verandi.vaga (conta_id, serie_id, pessoa_id, inicio, contrato_id) values ('
            f"'{conta_id}', '{serie['id']}', '{pessoa['id']}', '{inicio}', "
            f"{cita(contratos.get(pessoa['id']))}::uuid);")
    # as sessões que a agenda já tinha gerado não voltam a ser materializadas,
    # então a chamada delas recebe o aluno aqui, do mesmo jeito que o
    # materializador faria
    comandos.append(
        'insert into app_verandi.participacao (conta_id, sessao_id, pessoa_id, origem, status, '
        'registrado_por_origem) '
        "select s.conta_id, s.id, v.pessoa_id, 'recorrente', 'esperada', 'importacao' "
        'from app_verandi.sessao s join app_verandi.vaga v on v.serie_id = s.serie_id '
        "join app_verandi.conta c on c.id = s.conta_id "
        f"where s.conta_id = '{conta_id}' and s.status = 'prevista' "
        "and (s.inicio at time zone c.fuso)::date >= v.inicio "
        "on conflict (sessao_id, pessoa_id) do nothing;")
    comandos.append('commit;')
    sql('\n'.join(comandos))

    n = sql(f"select (select count(*) from app_verandi.vaga where conta_id = '{conta_id}')::int as v, "
            f"(select count(*) from app_verandi.participacao where conta_id = '{conta_id}')::int as p")[0]
    print(f"\ngravado: {n['v']} vagas, {n['p']} participações na conta")


if __name__ == '__main__':
    main()
